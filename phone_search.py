import csv
import glob
import json
import os
import re
import subprocess
import sys
import tkinter as tk
from datetime import datetime
from tkinter import filedialog, messagebox, ttk

# Optional Excel support
try:
    import openpyxl

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# Enable High DPI awareness on Windows
try:
    from ctypes import windll

    windll.shcore.SetProcessDpiAwareness(1)
except Exception:
    pass

# Brand synonyms mapping (bilingual support)
BRAND_SYNONYMS = {
    "xiaomi": ["xiaomi", "redmi", "mi", "小米", "红米"],
    "redmi": ["redmi", "xiaomi", "红米", "小米"],
    "apple": ["apple", "iphone", "ipad", "苹果"],
    "huawei": ["huawei", "华为"],
    "honor": ["honor", "荣耀"],
    "oneplus": ["oneplus", "1+", "一加"],
    "realme": ["realme", "真我"],
    "nubia": ["nubia", "redmagic", "努比亚", "红魔"],
    "google": ["google", "pixel", "谷歌"],
    "motorola": ["motorola", "moto", "摩托"],
    "nokia": ["nokia", "诺基亚"],
    "samsung": ["samsung", "galaxy", "三星"],
    "meizu": ["meizu", "魅族", "魅蓝"],
    "gionee": ["gionee", "金立"],
    "coolpad": ["coolpad", "ivvi", "酷派"],
    "zte": ["zte", "axon", "blade", "中兴", "远航", "天机"],
    "imoo": ["imoo", "步步高"],
    "hisense": ["hisense", "海信"],
    "asus": ["asus", "rog", "zenfone", "华硕"],
    "meitu": ["meitu", "美图"],
    "gree": ["gree", "格力"],
    "smartisan": ["smartisan", "锤子", "坚果"],
    "sugar": ["sugar", "糖果"],
    "gome": ["gome", "国美"],
    "lenovo": ["lenovo", "拯救者", "zuk", "联想"],
    "letv": ["letv", "leeco", "乐视"],
    "chinamobile": ["chinamobile", "移动", "中国移动"],
}


def clean_text(s):
    if not s:
        return ""
    return re.sub(r"[\s\-_+()（）/\[\]【】]+", "", str(s)).lower()


def get_base_dir():
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


class PhonePriceSearchApp:

    def __init__(self, root):
        self.root = root
        self.root.title("Phone Price Search Tool (Standalone Edition)")
        self.root.geometry("1180x700")
        self.root.minsize(880, 520)

        self.base_dir = get_base_dir()
        self.data_dir = os.path.join(self.base_dir, "data")
        self.history_file = os.path.join(self.data_dir, "price_history.json")

        if not os.path.exists(self.data_dir):
            os.makedirs(self.data_dir, exist_ok=True)

        self.all_data = []
        self.price_history = {}
        self.sort_column = None
        self.sort_reverse = False

        self._setup_ui()
        self.load_all_data()

    def _setup_ui(self):
        # 1. Top Search Bar
        top_frame = ttk.Frame(self.root, padding=(12, 10))
        top_frame.pack(fill=tk.X)

        ttk.Label(
            top_frame,
            text="Search (e.g. Xiaomi 8 / iPhone 13):",
            font=("Segoe UI", 10, "bold"),
        ).pack(side=tk.LEFT)

        self.search_var = tk.StringVar()
        self.search_var.trace_add("write", lambda *args: self.do_search())
        self.search_entry = ttk.Entry(
            top_frame, textvariable=self.search_var, font=("Segoe UI", 10), width=26
        )
        self.search_entry.pack(side=tk.LEFT, padx=6)

        clear_btn = ttk.Button(
            top_frame, text="Clear", command=lambda: self.search_var.set("")
        )
        clear_btn.pack(side=tk.LEFT, padx=3)

        self.only_changed_var = tk.BooleanVar(value=False)
        self.only_changed_var.trace_add("write", lambda *args: self.do_search())
        ttk.Checkbutton(
            top_frame,
            text="Show Price Changes Only",
            variable=self.only_changed_var,
        ).pack(side=tk.LEFT, padx=12)

        # Right Action Buttons
        btn_box = ttk.Frame(top_frame)
        btn_box.pack(side=tk.RIGHT)

        ttk.Button(
            btn_box, text="Batch Inventory Match", command=self.batch_price_inventory
        ).pack(side=tk.LEFT, padx=3)
        ttk.Button(
            btn_box, text="Open data/ Folder", command=self.open_data_folder
        ).pack(side=tk.LEFT, padx=3)
        ttk.Button(
            btn_box, text="Refresh", command=self.load_all_data
        ).pack(side=tk.LEFT, padx=3)

        # 2. Main Data Table
        table_frame = ttk.Frame(self.root, padding=(10, 0, 10, 5))
        table_frame.pack(fill=tk.BOTH, expand=True)

        self.columns = [
            ("source", "Brand / File", 110),
            ("series", "Series", 90),
            ("model", "Model / Edition", 210),
            ("price_good", "Screen Good", 100),
            ("price_screen_bad", "Screen Broken", 95),
            ("price_no_power", "No Power", 85),
            ("price_junk", "Scrap / Board", 85),
            ("fluctuation", "Price Trend", 140),
            ("remark", "Remarks (Double-click)", 260),
        ]

        self.tree = ttk.Treeview(
            table_frame,
            columns=[col[0] for col in self.columns],
            show="headings",
            selectmode="browse",
        )

        for col_id, col_name, width in self.columns:
            self.tree.heading(
                col_id,
                text=col_name,
                command=lambda c=col_id: self.sort_by_column(c),
            )
            self.tree.column(col_id, width=width, anchor=tk.W)

        self.tree.tag_configure("price_up", foreground="#d90429")
        self.tree.tag_configure("price_down", foreground="#2a9d8f")

        vsb = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=self.tree.yview)
        hsb = ttk.Scrollbar(table_frame, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")

        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)

        self.tree.bind("<Double-1>", self.show_detail)

        # 3. Status Bar
        self.status_var = tk.StringVar(value="Ready")
        status_bar = ttk.Label(
            self.root,
            textvariable=self.status_var,
            relief=tk.SUNKEN,
            anchor=tk.W,
            padding=(8, 4),
            font=("Segoe UI", 9),
        )
        status_bar.pack(fill=tk.X, side=tk.BOTTOM)

    def open_data_folder(self):
        if not os.path.exists(self.data_dir):
            os.makedirs(self.data_dir, exist_ok=True)
        if sys.platform == "win32":
            os.startfile(self.data_dir)
        elif sys.platform == "darwin":
            subprocess.Popen(["open", self.data_dir])
        else:
            subprocess.Popen(["xdg-open", self.data_dir])

    def load_all_data(self):
        self.all_data.clear()
        self.price_history.clear()

        if os.path.exists(self.history_file):
            try:
                with open(self.history_file, "r", encoding="utf-8") as f:
                    self.price_history = json.load(f)
            except Exception:
                self.price_history = {}

        csv_files = glob.glob(os.path.join(self.data_dir, "*.csv"))
        if not csv_files:
            for item in self.tree.get_children():
                self.tree.delete(item)
            self.status_var.set(
                f"Warning: No CSV files found in '{self.data_dir}'. Please add CSV files and click Refresh."
            )
            return

        encodings = ["utf-8-sig", "utf-8", "gb18030", "gbk"]
        loaded_files = 0

        for file_path in csv_files:
            file_name = (
                os.path.basename(file_path)
                .replace(".csv", "")
                .replace("_recycle_prices", "")
                .replace("回收价格", "")
            )

            file_content = None
            for enc in encodings:
                try:
                    with open(file_path, "r", encoding=enc, errors="strict") as f:
                        file_content = list(csv.DictReader(f))
                        break
                except Exception:
                    continue

            if not file_content:
                continue

            for row in file_content:
                model = (
                    row.get("型号")
                    or row.get("Model")
                    or row.get("品名/型号")
                    or row.get("Google")
                    or row.get("诺基亚")
                    or row.get("点数")
                    or row.get("点数系列")
                    or "Unknown"
                ).strip()
                series = (
                    row.get("系列")
                    or row.get("Series")
                    or row.get("品牌/系列")
                    or row.get("分类")
                    or file_name
                ).strip()
                p_good = (
                    row.get("开机屏好")
                    or row.get("开机靓好")
                    or row.get("开机屏好(无ID)")
                    or row.get("回收价格(元/点数)")
                    or "-"
                ).strip()
                p_s_bad = (
                    row.get("开机屏坏")
                    or row.get("开机屏坏(无ID)")
                    or row.get("开机屏坏/外碎")
                    or "-"
                ).strip()
                p_no_pwr = (row.get("不开机") or "-").strip()
                p_junk = (
                    row.get("废板-整机") or row.get("深板·整机") or "-"
                ).strip()
                remark = (row.get("备注") or row.get("Remarks") or "").strip()

                key = f"{file_name}_{model}"
                history_info = self.price_history.get(key)
                fluctuation = "-"
                change_tag = ""

                if history_info:
                    old_p = history_info.get("old_price", "")
                    diff = history_info.get("diff", 0)
                    if diff > 0:
                        fluctuation = f"UP +{diff} (Old: {old_p})"
                        change_tag = "price_up"
                    elif diff < 0:
                        fluctuation = f"DOWN -{abs(diff)} (Old: {old_p})"
                        change_tag = "price_down"

                self.all_data.append(
                    {
                        "source": file_name,
                        "series": series,
                        "model": model,
                        "price_good": p_good,
                        "price_screen_bad": p_s_bad,
                        "price_no_power": p_no_pwr,
                        "price_junk": p_junk,
                        "fluctuation": fluctuation,
                        "change_tag": change_tag,
                        "remark": remark,
                        "raw_dict": row,
                    }
                )
            loaded_files += 1

        changed_count = sum(1 for x in self.all_data if x["change_tag"])
        self.status_var.set(
            f"Loaded {loaded_files} files, {len(self.all_data)} models total. {changed_count} models have recent price updates."
        )
        self.do_search()

    def do_search(self):
        raw_query = self.search_var.get().strip()
        only_changed = self.only_changed_var.get()

        for item in self.tree.get_children():
            self.tree.delete(item)

        brand_terms, model_terms = self._parse_query(raw_query)
        matched_results = []

        for item in self.all_data:
            if only_changed and not item["change_tag"]:
                continue

            if not raw_query:
                matched_results.append((0, item))
                continue

            c_model = clean_text(item["model"])
            c_series = clean_text(item["series"])
            c_source = clean_text(item["source"])
            c_remark = clean_text(item["remark"])
            c_brand_scope = f"{c_source} {c_series}"

            # Brand scope filter
            if brand_terms:
                brand_matched = False
                for _, syn_list in brand_terms:
                    if any(
                        syn in c_brand_scope or c_model.startswith(syn)
                        for syn in syn_list
                    ):
                        brand_matched = True
                        break
                if not brand_matched:
                    continue

            # Model & variant matching
            model_matched = True
            score = 0
            for m_term in model_terms:
                if m_term in c_model:
                    if c_model == m_term:
                        score += 300
                    elif c_model.startswith(m_term) or c_model.endswith(m_term):
                        score += 180
                    else:
                        score += 100
                elif m_term in c_series or m_term in c_remark:
                    score += 40
                else:
                    model_matched = False
                    break

            if model_matched:
                score -= len(item["model"]) * 2
                matched_results.append((score, item))

        matched_results.sort(key=lambda x: x[0], reverse=True)

        for _, item in matched_results:
            tags = (item["change_tag"],) if item["change_tag"] else ()
            self.tree.insert(
                "",
                tk.END,
                values=(
                    item["source"],
                    item["series"],
                    item["model"],
                    item["price_good"],
                    item["price_screen_bad"],
                    item["price_no_power"],
                    item["price_junk"],
                    item["fluctuation"],
                    item["remark"],
                ),
                tags=tags,
            )

    def _parse_query(self, query_str):
        raw_terms = [t for t in re.split(r"[\s,+，]+", query_str.strip()) if t]
        if not raw_terms:
            return [], []

        if len(raw_terms) == 1:
            single = raw_terms[0]
            for brand, syns in BRAND_SYNONYMS.items():
                for s in syns:
                    if single.lower().startswith(s) and len(single) > len(s):
                        raw_terms = [single[: len(s)], single[len(s) :]]
                        break

        brand_terms = []
        model_terms = []
        for t in raw_terms:
            t_clean = clean_text(t)
            is_brand = False
            for brand_key, syns in BRAND_SYNONYMS.items():
                clean_syns = [clean_text(x) for x in syns]
                if t_clean in clean_syns:
                    brand_terms.append((t_clean, clean_syns))
                    is_brand = True
                    break
            if not is_brand:
                model_terms.append(t_clean)
        return brand_terms, model_terms

    def batch_price_inventory(self):
        file_path = filedialog.askopenfilename(
            title="Select Used Phone Inventory File",
            filetypes=[
                ("Spreadsheet Files", "*.xlsx *.xls *.csv"),
                ("Excel Files", "*.xlsx *.xls"),
                ("CSV Files", "*.csv"),
            ],
        )
        if not file_path:
            return

        rows = []
        if file_path.lower().endswith((".xlsx", ".xls")):
            if not HAS_OPENPYXL:
                messagebox.showerror(
                    "Missing openpyxl",
                    "To read Excel files (.xlsx), please run:\npip install openpyxl\nOr save your spreadsheet as CSV and import again.",
                )
                return
            wb = openpyxl.load_workbook(file_path)
            ws = wb.active
            headers = [str(cell.value or "").strip() for cell in ws[1]]
            for r in ws.iter_rows(min_row=2, values_only=True):
                rows.append(
                    {
                        headers[i]: str(r[i] or "").strip()
                        for i in range(len(headers))
                        if i < len(r)
                    }
                )
        else:
            encodings = ["utf-8-sig", "gb18030", "utf-8", "gbk"]
            for enc in encodings:
                try:
                    with open(file_path, "r", encoding=enc) as f:
                        rows = list(csv.DictReader(f))
                        break
                except Exception:
                    continue

        if not rows:
            messagebox.showwarning("Warning", "No data rows found in the selected file.")
            return

        sample_keys = list(rows[0].keys())
        model_col = next(
            (
                k
                for k in sample_keys
                if any(x in k.lower() for x in ["model", "device", "name", "型号", "机型", "品名"])
            ),
            sample_keys[0],
        )

        matched_count = 0
        result_rows = []

        for r in rows:
            raw_input_model = r.get(model_col, "").strip()
            c_input = clean_text(raw_input_model)

            best_match = None
            highest_score = -1

            for db_item in self.all_data:
                db_model = clean_text(db_item["model"])
                db_source = clean_text(db_item["source"])

                if c_input == db_model or c_input == f"{db_source}{db_model}":
                    best_match = db_item
                    highest_score = 1000
                    break

                if db_model and (db_model in c_input or c_input in db_model):
                    score = 500 - abs(len(c_input) - len(db_model)) * 10
                    if score > highest_score:
                        highest_score = score
                        best_match = db_item

            row_out = dict(r)
            if best_match:
                row_out["Matched_Model"] = f"[{best_match['source']}] {best_match['model']}"
                row_out["Quote_ScreenGood"] = best_match["price_good"]
                row_out["Quote_ScreenBroken"] = best_match["price_screen_bad"]
                row_out["Quote_NoPower"] = best_match["price_no_power"]
                row_out["Quote_ScrapBoard"] = best_match["price_junk"]
                row_out["Deduction_Remarks"] = best_match["remark"]
                matched_count += 1
            else:
                row_out["Matched_Model"] = "Unmatched"
                row_out["Quote_ScreenGood"] = "-"
                row_out["Quote_ScreenBroken"] = "-"
                row_out["Quote_NoPower"] = "-"
                row_out["Quote_ScrapBoard"] = "-"
                row_out["Deduction_Remarks"] = "-"

            result_rows.append(row_out)

        default_out_name = f"Priced_{os.path.splitext(os.path.basename(file_path))[0]}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        save_path = filedialog.asksaveasfilename(
            title="Save Priced Inventory",
            initialfile=default_out_name,
            defaultextension=".xlsx" if HAS_OPENPYXL else ".csv",
            filetypes=[("Excel Files", "*.xlsx")] if HAS_OPENPYXL else [("CSV Files", "*.csv")],
        )
        if not save_path:
            return

        if save_path.lower().endswith(".xlsx") and HAS_OPENPYXL:
            wb_out = openpyxl.Workbook()
            ws_out = wb_out.active
            ws_out.title = "Appraisal_Result"
            out_headers = list(result_rows[0].keys())
            ws_out.append(out_headers)
            for item in result_rows:
                ws_out.append([item.get(h, "") for h in out_headers])
            wb_out.save(save_path)
        else:
            with open(save_path, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=list(result_rows[0].keys()))
                writer.writeheader()
                writer.writerows(result_rows)

        messagebox.showinfo(
            "Completed",
            f"Batch pricing completed!\nTotal rows: {len(rows)}\nMatched models: {matched_count}\n\nSaved to:\n{save_path}",
        )

    def sort_by_column(self, col):
        if self.sort_column == col:
            self.sort_reverse = not self.sort_reverse
        else:
            self.sort_reverse = False
            self.sort_column = col

        def sort_key(item):
            val = item.get(col, "")
            try:
                return (0, float(val))
            except ValueError:
                return (1, str(val))

        self.all_data.sort(key=sort_key, reverse=self.sort_reverse)
        self.do_search()

    def show_detail(self, event):
        selected = self.tree.selection()
        if not selected:
            return

        values = self.tree.item(selected[0], "values")
        source_name, model_name = values[0], values[2]

        matched_item = next(
            (
                item
                for item in self.all_data
                if item["source"] == source_name and item["model"] == model_name
            ),
            None,
        )

        if not matched_item:
            return

        win = tk.Toplevel(self.root)
        win.title(f"Price & Deduction Details - {model_name}")
        win.geometry("560x460")
        win.transient(self.root)

        frame = ttk.Frame(win, padding=15)
        frame.pack(fill=tk.BOTH, expand=True)

        ttk.Label(
            frame,
            text=f"[{source_name}] {model_name}",
            font=("Segoe UI", 12, "bold"),
        ).pack(anchor=tk.W, pady=(0, 10))

        text_box = tk.Text(
            frame, wrap=tk.WORD, font=("Consolas", 10), padx=8, pady=8
        )
        text_box.pack(fill=tk.BOTH, expand=True)

        for k, v in matched_item["raw_dict"].items():
            if v and str(v).strip():
                text_box.insert(tk.END, f"- {k.ljust(18)} : {v}\n")

        if matched_item.get("change_tag"):
            text_box.insert(
                tk.END,
                f"\nPrice Trend History: {matched_item['fluctuation']}\n",
            )

        text_box.configure(state="disabled")
        ttk.Button(frame, text="Close", command=win.destroy).pack(pady=10)


if __name__ == "__main__":
    root = tk.Tk()
    app = PhonePriceSearchApp(root)
    root.mainloop()
